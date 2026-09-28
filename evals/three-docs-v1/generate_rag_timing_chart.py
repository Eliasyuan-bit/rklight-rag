"""Generate an editable Excel comparison of Naive and Graph RAG latency."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT = Path(__file__).resolve().parents[2] / "docs" / "rag-naive-graph-timing-comparison.xlsx"


def style_header(ws, row, first_col, last_col):
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=first_col, max_col=last_col):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def style_table(ws, first_row, last_row, first_col, last_col):
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="243447")
            cell.alignment = Alignment(vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.0'
                cell.alignment = Alignment(horizontal="right", vertical="center")


wb = Workbook()
overview = wb.active
overview.title = "耗时对比图"
data = wb.create_sheet("数据与口径")
wb.calculation.fullCalcOnLoad = True

overview.merge_cells("B2:P2")
overview["B2"] = "Naive RAG 与 Graph RAG · 问答耗时对比"
overview["B2"].font = Font(name="Microsoft YaHei", size=20, bold=True, color="17365D")
overview.merge_cells("B3:P3")
overview["B3"] = "测试问题：FT 测试中 USB 错误是为什么？  ·  单位：毫秒（ms）"
overview["B3"].font = Font(name="Microsoft YaHei", size=11, color="58708B")
overview.merge_cells("B4:P4")
overview["B4"] = "上图为阶段耗时的横向堆叠比较，不是按真实开始时间绘制的甘特图；灰色为阶段合计与端到端总耗时的差额。"
overview["B4"].font = Font(name="Microsoft YaHei", size=10, color="A04432")
overview.merge_cells("B5:P5")
overview["B5"] = "注意：Rerank 耗时取后一次复测，与其他阶段及端到端耗时不是同一请求；请勿将灰色部分解释为实际发生在回答之后。"
overview["B5"].font = Font(name="Microsoft YaHei", size=10, color="A04432")
overview.column_dimensions["A"].width = 3
for col in range(2, 17):
    overview.column_dimensions[get_column_letter(col)].width = 12
overview.sheet_view.showGridLines = False
overview.freeze_panes = "B6"

data.append(["阶段耗时与端到端差额（ms）"])
data.append(["方案", "查询理解 LLM", "Embedding", "Rerank", "回答 LLM", "未归因差额", "端到端总耗时", "阶段合计"])
data.append(["Graph RAG", 976.4, 92.6, 681.4, 2230.3, 879.4, 4860.1, 3980.7])
data.append(["Naive RAG", 0, 30.6, 226.1, 2798.4, 744.1, 3799.2, 3055.1])
style_header(data, 2, 1, 8)
style_table(data, 3, 4, 1, 8)

data["A7"] = "端到端指标（ms）"
data.append([])
data["A8"] = "指标"
data["B8"] = "Graph RAG"
data["C8"] = "Naive RAG"
data["A9"] = "系统首字时间 TTFT"
data["B9"] = 3680.1
data["C9"] = 2277.6
data["A10"] = "总耗时"
data["B10"] = 4860.1
data["C10"] = 3799.2
style_header(data, 8, 1, 3)
style_table(data, 9, 10, 1, 3)

data["A13"] = "输入规模（用于解释耗时差异）"
for col, value in enumerate(["阶段", "Naive RAG", "Graph RAG"], start=1):
    data.cell(14, col, value)
for row, values in enumerate([
    ("查询理解 LLM", "不调用", "662 输入 / 41 输出 token"),
    ("Embedding", "1 条 / 11 输入 token", "3 条 / 37 输入 token"),
    ("Rerank", "1 个候选 / 762 输入 token", "6 个候选 / 2340 输入 token"),
    ("回答 LLM", "1613 输入 / 142 输出 token", "1188 输入 / 117 输出 token"),
    ("全部 LLM", "1613 输入 / 142 输出 token", "1850 输入 / 158 输出 token"),
], start=15):
    for col, value in enumerate(values, start=1):
        data.cell(row, col, value)
style_header(data, 14, 1, 3)
style_table(data, 15, 19, 1, 3)

data["A22"] = "口径与限制"
notes = [
    "堆叠图将阶段耗时按逻辑顺序排布，横轴仅表示耗时长度，不代表真实开始/结束时间。",
    "未归因差额 = 端到端总耗时 - 阶段合计；Naive 为 744.1 ms，Graph 为 879.4 ms。",
    "Rerank 的 226.1 / 681.4 ms 取后一次 token 计数复测；其余阶段和端到端数据取原测试。",
    "因此图中的未归因差额只是跨测试数据的算术差值，不应解释为精确的框架开销。",
    "Graph RAG 在本项目中对应 LightRAG mix 检索模式。以上均为单次样本，非稳态统计值。",
]
for row, note in enumerate(notes, start=23):
    data.cell(row, 1, f"• {note}")
    data.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    data.cell(row, 1).font = Font(name="Microsoft YaHei", size=10, color="52677D")
data.column_dimensions["A"].width = 28
for col in "BCDEFGH":
    data.column_dimensions[col].width = 23
data.column_dimensions["C"].width = 29
data.sheet_view.showGridLines = False
data.freeze_panes = "B3"

stage_chart = BarChart()
stage_chart.type = "bar"
stage_chart.grouping = "stacked"
stage_chart.overlap = 100
stage_chart.gapWidth = 90
stage_chart.title = "阶段耗时与端到端差额"
stage_chart.x_axis.title = "耗时（ms）"
stage_chart.y_axis.title = "方案"
stage_chart.width = 25
stage_chart.height = 9.5
stage_chart.style = 10
stage_chart.legend.position = "b"
stage_chart.x_axis.scaling.min = 0
stage_chart.x_axis.scaling.max = 5200
stage_chart.x_axis.majorUnit = 1000
stage_chart.add_data(Reference(data, min_col=2, max_col=6, min_row=2, max_row=4), titles_from_data=True)
stage_chart.set_categories(Reference(data, min_col=1, min_row=3, max_row=4))
for series, color in zip(stage_chart.series, ["7353BA", "32B6A0", "EF9D41", "326FA8", "B7C2CC"]):
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line = GraphicalProperties(noFill=True).line
overview.add_chart(stage_chart, "B7")

metric_chart = BarChart()
metric_chart.type = "bar"
metric_chart.grouping = "clustered"
metric_chart.title = "端到端：首字时间与总耗时"
metric_chart.x_axis.title = "耗时（ms）"
metric_chart.width = 25
metric_chart.height = 8
metric_chart.style = 10
metric_chart.legend.position = "b"
metric_chart.x_axis.scaling.min = 0
metric_chart.x_axis.scaling.max = 5200
metric_chart.x_axis.majorUnit = 1000
metric_chart.add_data(Reference(data, min_col=2, max_col=3, min_row=8, max_row=10), titles_from_data=True)
metric_chart.set_categories(Reference(data, min_col=1, min_row=9, max_row=10))
for series, color in zip(metric_chart.series, ["7353BA", "326FA8"]):
    series.graphicalProperties.solidFill = color
metric_chart.dataLabels = DataLabelList()
metric_chart.dataLabels.showVal = True
overview.add_chart(metric_chart, "B27")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
