#!/usr/bin/env python3
"""Build the warm-state Excel and HTML comparison from one measured JSON file."""

import html
import json
from pathlib import Path
from statistics import median

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).resolve().parent / "results/qa-usb-warm-uncached-timing-20260921.json"
XLSX = ROOT / "docs/rag-naive-vs-mix-warm-performance.xlsx"
HTML = ROOT / "docs/rag-naive-vs-graph-timing.html"


def other(row):
    return row["total_ms"] - row["keyword_llm_ms"] - row["embedding_ms"] - row["rerank_ms"] - row["answer_llm_ms"]


def summary(rows):
    return {
        "ttft": round(median(row["ttft_ms"] for row in rows), 1),
        "total": round(median(row["total_ms"] for row in rows), 1),
        "keyword_llm": round(median(row["keyword_llm_ms"] for row in rows), 1),
        "embedding": round(median(row["embedding_ms"] for row in rows), 1),
        "rerank": round(median(row["rerank_ms"] for row in rows), 1),
        "answer_llm": round(median(row["answer_llm_ms"] for row in rows), 1),
        "other": round(median(other(row) for row in rows), 1),
    }


def build_excel(data, naive, mix):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "稳定态对比"
    raw = wb.create_sheet("逐次记录")
    sheet.sheet_view.showGridLines = False
    raw.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Naive RAG 与 LightRAG Mix：稳定态性能"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color="17365D")
    sheet.merge_cells("A2:G2")
    sheet["A2"] = "同一 USB 问题；Naive 与 Mix 各预热 1 次后测试 5 次；单客户端；Rerank 开启。数值为逐次请求的中位数。"
    sheet.merge_cells("A3:G3")
    sheet["A3"] = "LLM 缓存已临时关闭：Mix 每次均执行查询理解和回答两次 LLM；测试结束后恢复默认缓存设置。"
    sheet["A3"].font = Font(name="Microsoft YaHei", color="A04432", bold=True, size=10)
    sheet.merge_cells("A4:G4")
    sheet["A4"] = "横向堆叠图按逻辑阶段排布，不代表真实时间戳；各阶段分别取中位数，不能直接相加得到总耗时。"
    sheet["A4"].font = Font(name="Microsoft YaHei", color="6D7785", size=10)
    sheet.merge_cells("A5:G5")
    sheet["A5"] = "“非模型耗时”是每条请求先扣除全部模型调用、再取中位数；它分散发生在检索、调度和传输等位置。"
    sheet["A5"].font = Font(name="Microsoft YaHei", color="6D7785", size=10)

    sheet.append([])
    sheet.append(["指标", "Naive RAG", "LightRAG Mix", "Mix - Naive"])
    metrics = [
        ("系统首字时间 TTFT (ms)", "ttft"),
        ("端到端总耗时 (ms)", "total"),
        ("查询理解 LLM (ms)", "keyword_llm"),
        ("Embedding (ms)", "embedding"),
        ("Rerank (ms)", "rerank"),
        ("回答 LLM (ms)", "answer_llm"),
        ("非模型耗时：检索/调度/传输等 (ms)", "other"),
    ]
    for label, key in metrics:
        sheet.append([label, naive[key], mix[key], round(mix[key] - naive[key], 1)])
    for c in sheet[6][:4]:
        c.fill = PatternFill("solid", fgColor="17365D")
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for row in sheet.iter_rows(min_row=7, max_row=13, min_col=2, max_col=4):
        for c in row:
            c.number_format = '#,##0.0'
            c.alignment = Alignment(horizontal="right")

    sheet["A16"] = "输入规模（每次相同）"
    for c, value in zip(sheet[17][:3], ["阶段", "Naive RAG", "LightRAG Mix"]):
        c.value = value
        c.fill = PatternFill("solid", fgColor="17365D")
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    rows = [
        ("查询理解 LLM", "不调用", "662 输入 / 41 输出 token"),
        ("Embedding", "1 条 / 12 输入 token", "3 条 / 34 输入 token"),
        ("Rerank", "1 候选 / 762 输入 token", "6 候选 / 2340 输入 token"),
        ("回答 LLM", "1613 输入 / 142 输出 token", "1188 输入 / 117 输出 token"),
        ("有效样本", "预热后 5 次", "预热后 5 次；LLM 缓存关闭"),
    ]
    for r, values in enumerate(rows, start=18):
        for c, value in enumerate(values, start=1):
            sheet.cell(r, c, value)

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 18
    sheet.freeze_panes = "B7"

    sheet["A32"] = "堆叠图数据：分别取中位数"
    gantt_headers = ["方案", "查询理解 LLM", "Embedding", "Rerank", "非模型耗时（汇总）", "回答 LLM"]
    for col, label in enumerate(gantt_headers, start=1):
        sheet.cell(33, col, label)
    for row, mode, values in [(34, "LightRAG Mix", mix), (35, "Naive RAG", naive)]:
        for col, value in enumerate([
            mode, values["keyword_llm"], values["embedding"], values["rerank"],
            values["other"], values["answer_llm"],
        ], start=1):
            sheet.cell(row, col, value)
    sheet["A37"] = "阶段中位数之和"
    sheet["B37"] = round(sum(naive[key] for key in ("keyword_llm", "embedding", "rerank", "other", "answer_llm")), 1)
    sheet["C37"] = round(sum(mix[key] for key in ("keyword_llm", "embedding", "rerank", "other", "answer_llm")), 1)
    sheet["A38"] = "端到端总耗时中位数"
    sheet["B38"] = naive["total"]
    sheet["C38"] = mix["total"]
    sheet["A39"] = "两种中位数口径的差额"
    sheet["B39"] = round(naive["total"] - sheet["B37"].value, 1)
    sheet["C39"] = round(mix["total"] - sheet["C37"].value, 1)

    gantt = BarChart()
    gantt.type = "bar"
    gantt.grouping = "stacked"
    gantt.overlap = 100
    gantt.gapWidth = 85
    gantt.title = "各阶段耗时中位数 · 横向堆叠图"
    gantt.x_axis.title = "耗时（ms）"
    gantt.width = 27
    gantt.height = 9
    gantt.legend.position = "b"
    gantt.add_data(Reference(sheet, min_col=2, max_col=6, min_row=33, max_row=35), titles_from_data=True)
    gantt.set_categories(Reference(sheet, min_col=1, min_row=34, max_row=35))
    for series, color in zip(gantt.series, ["7353BA", "32B6A0", "EF9D41", "AEB9C6", "326FA8"]):
        series.graphicalProperties.solidFill = color
    sheet.add_chart(gantt, "F6")

    stage = BarChart()
    stage.type = "bar"
    stage.grouping = "clustered"
    stage.title = "阶段耗时中位数（不是连续时间线）"
    stage.x_axis.title = "毫秒"
    stage.width = 24
    stage.height = 10
    stage.legend.position = "b"
    # Build from a compact backing table to keep chart categories contiguous.
    sheet["A24"] = "阶段"
    sheet["B24"] = "Naive RAG"
    sheet["C24"] = "LightRAG Mix"
    for r, (label, key) in enumerate(metrics[2:], start=25):
        sheet.cell(r, 1, label)
        sheet.cell(r, 2, naive[key])
        sheet.cell(r, 3, mix[key])
    stage.add_data(Reference(sheet, min_col=2, max_col=3, min_row=24, max_row=29), titles_from_data=True)
    stage.set_categories(Reference(sheet, min_col=1, min_row=25, max_row=29))
    stage.series[0].graphicalProperties.solidFill = "326FA8"
    stage.series[1].graphicalProperties.solidFill = "EF9D41"
    stage.dataLabels = DataLabelList()
    stage.dataLabels.showVal = True
    sheet.add_chart(stage, "F27")

    e2e = BarChart()
    e2e.type = "bar"
    e2e.grouping = "clustered"
    e2e.title = "端到端耗时中位数"
    e2e.x_axis.title = "毫秒"
    e2e.width = 24
    e2e.height = 7
    e2e.legend.position = "b"
    e2e.add_data(Reference(sheet, min_col=2, max_col=3, min_row=6, max_row=8), titles_from_data=True)
    e2e.set_categories(Reference(sheet, min_col=1, min_row=7, max_row=8))
    for series, color in zip(e2e.series, ["326FA8", "EF9D41"]):
        series.graphicalProperties.solidFill = color
    e2e.dataLabels = DataLabelList()
    e2e.dataLabels.showVal = True
    sheet.add_chart(e2e, "F48")

    raw.append(["模式", "轮次", "TTFT (ms)", "总耗时 (ms)", "查询理解 LLM (ms)", "Embedding (ms)", "Rerank (ms)", "回答 LLM (ms)", "非模型耗时 (ms)"])
    for mode, items in [("Naive RAG", data["naive"]), ("LightRAG Mix", data["mix"])]:
        for row in items:
            raw.append([mode, row["run"], row["ttft_ms"], row["total_ms"], row["keyword_llm_ms"], row["embedding_ms"], row["rerank_ms"], row["answer_llm_ms"], round(other(row), 1)])
    for c in raw[1]:
        c.fill = PatternFill("solid", fgColor="17365D")
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for column in "ABCDEFGHI":
        raw.column_dimensions[column].width = 23
    raw.freeze_panes = "C2"
    wb.save(XLSX)


def build_html(naive, mix):
    items = [
        ("系统首字时间 TTFT", "ttft", 4300),
        ("端到端总耗时", "total", 4300),
        ("查询理解 LLM", "keyword_llm", 2800),
        ("Embedding", "embedding", 2800),
        ("Rerank", "rerank", 2800),
        ("回答 LLM", "answer_llm", 2800),
        ("非模型耗时", "other", 2800),
    ]
    groups = []
    for label, key, scale in items:
        bars = []
        for mode, values, css in [("Naive RAG", naive, "naive"), ("LightRAG Mix", mix, "mix")]:
            value = values[key]
            bars.append(
                f'<div class="row"><span>{mode}</span><div class="track"><div class="bar {css}" '
                f'style="width:{100 * value / scale:.2f}%"></div></div><strong>{value:.1f} ms</strong></div>'
            )
        groups.append(f'<div class="metric"><h2>{html.escape(label)}</h2>{"".join(bars)}</div>')
    page = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Naive RAG 与 LightRAG Mix 稳定态耗时</title>
<style>
body{{font-family:system-ui,"Microsoft YaHei",sans-serif;background:#f4f7fb;color:#23334a;margin:0;padding:32px}}
main{{max-width:1080px;margin:auto}}h1{{font-size:26px;margin:0 0 8px}}p{{line-height:1.6}}
.lead{{color:#60748c;margin:0 0 20px}}.note{{background:#fff5e9;border-left:4px solid #ef9d41;padding:12px 16px;border-radius:6px;margin:18px 0}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(420px,1fr));gap:16px}}.metric{{background:white;border:1px solid #dfe7f0;border-radius:12px;padding:16px 18px;box-shadow:0 3px 12px #17365d0a}}
h2{{font-size:16px;margin:0 0 14px}}.row{{display:grid;grid-template-columns:108px 1fr 96px;gap:10px;align-items:center;font-size:13px;margin:10px 0}}
.track{{height:20px;border-radius:5px;background:#eaf0f6;overflow:hidden}}.bar{{height:100%;border-radius:5px}}.naive{{background:#326fa8}}.mix{{background:#ef9d41}}strong{{text-align:right;font-variant-numeric:tabular-nums}}
.foot{{font-size:13px;color:#60748c;margin-top:20px}}a{{color:#326fa8}}
@media(max-width:600px){{body{{padding:16px}}.grid{{grid-template-columns:1fr}}.row{{grid-template-columns:92px 1fr 80px;font-size:11px}}}}
</style></head><body><main>
<h1>Naive RAG 与 LightRAG Mix 稳定态耗时</h1>
<p class="lead">同一问题“FT 测试中 USB 错误是为什么？”；Naive 与 Mix 各预热 1 次后测试 5 次；图示均为中位数。</p>
<div class="note"><b>测试口径：</b>临时关闭 LLM 缓存，Mix 每次都执行“查询理解 LLM + 回答 LLM”两次调用；测试后已恢复默认设置。Rerank 的热/冷态仍可能变化，逐次记录见 Excel。</div>
<div class="grid">{"".join(groups)}</div>
<p class="foot">非模型耗时＝每次端到端总耗时减去查询理解 LLM、Embedding、Rerank 与回答 LLM 的网关耗时，再取中位数；包含 LightRAG 检索处理、调度与传输等。各指标分别取中位数，不能横向相加。<a href="rag-naive-vs-mix-warm-performance.xlsx">下载 Excel</a>。</p>
</main></body></html>
'''
    HTML.write_text(page, encoding="utf-8")


def main():
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    naive, mix = summary(data["naive"]), summary(data["mix"])
    build_excel(data, naive, mix)
    build_html(naive, mix)
    print(XLSX)
    print(HTML)


if __name__ == "__main__":
    main()
